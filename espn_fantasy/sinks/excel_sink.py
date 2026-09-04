"""Excel: one workbook, one worksheet per table.

For the many leagues whose history ends up being argued about in a
spreadsheet. Frozen header rows and an auto-filter on every sheet, since that
is the first thing anyone does anyway.
"""

from __future__ import annotations

import logging

from ..errors import MissingDependency
from ..parse import Dataset
from . import Sink, SinkResult
from .files import as_text, out_dir

log = logging.getLogger(__name__)

#: Excel refuses to store a longer string in a single cell. `settings_raw`
#: routinely exceeds it, so it is truncated with a visible marker rather than
#: silently corrupting the workbook.
MAX_CELL = 32767
TRUNCATION_MARKER = "...[truncated]"


class ExcelSink(Sink):
    name = "xlsx"
    requires = ("openpyxl", "openpyxl")
    extra = "excel"
    help = "one .xlsx workbook, a sheet per table (needs [excel] extra)"

    def write(self, data: Dataset) -> SinkResult:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except ModuleNotFoundError as exc:
            raise MissingDependency("xlsx", "openpyxl", "excel") from exc

        target = out_dir(self.cfg)
        path = target / "espn_fantasy.xlsx"
        workbook = Workbook(write_only=False)
        workbook.remove(workbook.active)

        truncated = 0
        total = 0
        for table, rows in data:
            sheet = workbook.create_sheet(title=table.name[:31])
            sheet.append(list(table.columns))
            for row in rows:
                values = []
                for column in table.columns:
                    value = as_text(row.get(column), column, table)
                    if isinstance(value, str) and len(value) > MAX_CELL:
                        value = value[:MAX_CELL - len(TRUNCATION_MARKER)] + TRUNCATION_MARKER
                        truncated += 1
                    values.append(value)
                sheet.append(values)
            sheet.freeze_panes = "A2"
            if rows:
                last = get_column_letter(len(table.columns))
                sheet.auto_filter.ref = f"A1:{last}{len(rows) + 1}"
            total += len(rows)

        workbook.save(path)
        if truncated:
            log.warning(
                "%d cell(s) exceeded Excel's %d-character limit and were "
                "truncated; the full values are in any other output format",
                truncated, MAX_CELL)
        return SinkResult("xlsx", str(path), total)
